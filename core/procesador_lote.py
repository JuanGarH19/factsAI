"""
procesador_lote.py
------------------
FacturaIA — Procesador por lotes de facturas.

Escanea una carpeta (o lista de archivos), extrae los datos de cada factura
usando FacturaExtractor como motor, y exporta los resultados a Excel y/o CSV.

Uso básico
----------
    from procesador_lote import ProcesadorLote

    procesador = ProcesadorLote()
    informe = procesador.procesar_carpeta("/ruta/a/facturas/")
    informe.guardar_ambos("/home/juang/factura-ia/data/output/")

Uso por CLI
-----------
    python procesador_lote.py /ruta/carpeta/facturas
    python procesador_lote.py /ruta/carpeta/facturas --output /home/juang/factura-ia/data/output
    python procesador_lote.py /ruta/carpeta/facturas --solo-csv
    python procesador_lote.py /ruta/carpeta/facturas --solo-excel
    python procesador_lote.py factura1.pdf factura2.jpg   # archivos sueltos
"""

from __future__ import annotations

import csv
import logging
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .extractor import (
    EXTENSIONES_IMAGEN, EXTENSIONES_PDF,
    FacturaExtractor, ResultadoExtraccion,
)

logger = logging.getLogger(__name__)

# ─────────────────────────────────────────────
# Constantes de formato
# ─────────────────────────────────────────────

OUTPUT_DIR_DEFAULT = Path("/home/juang/factura-ia/data/output")

COLOR_HEADER_BG   = "1F4E79"
COLOR_VALIDA      = "E2EFDA"
COLOR_ADVERTENCIA = "FFF2CC"
COLOR_ERROR       = "FCE4D6"
COLOR_TOTAL_BG    = "D6DCE4"
COLOR_RESUMEN_BG  = "DEEAF1"

COLUMNAS_EXCEL = [
    ("archivo",          "Archivo",            22),
    ("nombre_emisor",    "Proveedor",           28),
    ("nif_emisor",       "NIF/CIF",             13),
    ("numero_factura",   "Nº Factura",          18),
    ("fecha_emision",    "Fecha Emisión",        14),
    ("fecha_vencimiento","Vencimiento",          14),
    ("concepto",         "Concepto",             35),
    ("base_imponible",   "Base Imponible (€)",  18),
    ("tipo_iva",         "IVA (%)",              10),
    ("cuota_iva",        "Cuota IVA (€)",        15),
    ("total_factura",    "Total Factura (€)",    18),
    ("metodo_pago",      "Forma de Pago",        16),
    ("iban",             "IBAN",                 28),
    ("moneda",           "Moneda",                9),
    ("_estado",          "Estado",               13),
    ("_confianza",       "Confianza",            11),
    ("error_log",        "Log de Errores",       40),
]

COLUMNAS_RESUMEN = [
    ("nif_emisor",    "NIF/CIF",             13),
    ("nombre_emisor", "Proveedor",           28),
    ("_n_facturas",   "Nº Facturas",         12),
    ("base_imponible","Base Imponible (€)",  18),
    ("cuota_iva",     "Cuota IVA (€)",       15),
    ("total_factura", "Total Factura (€)",   18),
]


# ─────────────────────────────────────────────
# Fila de resultado plano
# ─────────────────────────────────────────────

@dataclass
class FilaFactura:
    archivo: str
    datos: dict[str, Any]
    estado: str
    confianza: float
    advertencias: list[str]
    error_log: str = ""

    def valor(self, campo: str) -> Any:
        if campo == "archivo":    return self.archivo
        if campo == "_estado":    return self.estado
        if campo == "_confianza": return f"{self.confianza:.0%}"
        if campo == "error_log":  return self.error_log
        return self.datos.get(campo)


# ─────────────────────────────────────────────
# Informe de lote
# ─────────────────────────────────────────────

@dataclass
class InformeLote:
    filas: list[FilaFactura] = field(default_factory=list)
    timestamp: str = field(default_factory=lambda: datetime.now().strftime("%Y%m%d_%H%M%S"))

    @property
    def total(self) -> int:        return len(self.filas)
    @property
    def validas(self) -> int:      return sum(1 for f in self.filas if f.estado == "válida")
    @property
    def con_advertencias(self) -> int: return sum(1 for f in self.filas if f.estado == "incompleta")
    @property
    def con_error(self) -> int:    return sum(1 for f in self.filas if f.estado == "error")

    def guardar_excel(self, directorio: str | Path = OUTPUT_DIR_DEFAULT) -> Path:
        directorio = Path(directorio)
        directorio.mkdir(parents=True, exist_ok=True)
        ruta = directorio / f"facturas_{self.timestamp}.xlsx"
        self._construir_excel(ruta)
        logger.info(f"✅ Excel guardado: {ruta}")
        return ruta

    def guardar_csv(self, directorio: str | Path = OUTPUT_DIR_DEFAULT) -> Path:
        directorio = Path(directorio)
        directorio.mkdir(parents=True, exist_ok=True)
        ruta = directorio / f"facturas_{self.timestamp}.csv"
        self._construir_csv(ruta)
        logger.info(f"✅ CSV guardado: {ruta}")
        return ruta

    def guardar_ambos(self, directorio: str | Path = OUTPUT_DIR_DEFAULT) -> tuple[Path, Path]:
        return self.guardar_excel(directorio), self.guardar_csv(directorio)

    # ── Construcción Excel ────────────────────

    def _construir_excel(self, ruta: Path):
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Facturas"
        ws2 = wb.create_sheet("Resumen por Proveedor")
        self._hoja_facturas(ws1)
        self._hoja_resumen(ws2)
        ws1.freeze_panes = "A3"
        ws2.freeze_panes = "A3"
        wb.save(ruta)

    def _border(self):
        thin = Side(style="thin", color="CCCCCC")
        return Border(left=thin, right=thin, top=thin, bottom=thin)

    def _hoja_facturas(self, ws):
        border = self._border()
        font   = Font(name="Arial", size=10)
        bold   = Font(name="Arial", size=10, bold=True)

        # Fila 1: Título
        n_cols = len(COLUMNAS_EXCEL)
        ws.merge_cells(f"A1:{get_column_letter(n_cols)}1")
        t = ws["A1"]
        t.value     = f"FacturaIA — Informe de Facturas  ·  {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        t.font      = Font(name="Arial", size=12, bold=True, color="FFFFFF")
        t.fill      = PatternFill("solid", start_color="0D3B66")
        t.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        # Fila 2: Cabecera
        hdr_fill = PatternFill("solid", start_color=COLOR_HEADER_BG)
        hdr_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        for c, (campo, etiqueta, ancho) in enumerate(COLUMNAS_EXCEL, 1):
            cell = ws.cell(row=2, column=c, value=etiqueta)
            cell.font      = hdr_font
            cell.fill      = hdr_fill
            cell.border    = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(c)].width = ancho
        ws.row_dimensions[2].height = 30

        # Índices de columnas numéricas
        def col_of(nombre):
            return next(i for i, (c, *_) in enumerate(COLUMNAS_EXCEL, 1) if c == nombre)

        col_base  = col_of("base_imponible")
        col_cuota = col_of("cuota_iva")
        col_total = col_of("total_factura")
        col_n_fac = col_of("numero_factura")

        # Filas de datos (desde fila 3)
        for r, fila in enumerate(self.filas, 3):
            if fila.estado == "error":
                color = COLOR_ERROR
            elif fila.advertencias:
                color = COLOR_ADVERTENCIA
            else:
                color = COLOR_VALIDA if r % 2 == 0 else "FFFFFF"
            fill = PatternFill("solid", start_color=color)

            for c, (campo, _, _) in enumerate(COLUMNAS_EXCEL, 1):
                valor = fila.valor(campo)
                cell  = ws.cell(row=r, column=c, value=valor)
                cell.font   = font
                cell.fill   = fill
                cell.border = border

                if campo in ("base_imponible", "cuota_iva", "total_factura") and isinstance(valor, (int, float)):
                    cell.number_format = '#,##0.00 €'
                    cell.alignment     = Alignment(horizontal="right", vertical="center")
                elif campo == "tipo_iva" and isinstance(valor, (int, float)):
                    cell.number_format = '0"%"'
                    cell.alignment     = Alignment(horizontal="center", vertical="center")
                elif campo == "_confianza":
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif campo == "_estado":
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if valor == "válida":
                        cell.font = Font(name="Arial", size=10, bold=True, color="375623")
                    elif valor == "error":
                        cell.font = Font(name="Arial", size=10, bold=True, color="9C0006")
                    else:
                        cell.font = Font(name="Arial", size=10, color="7D6608")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
            ws.row_dimensions[r].height = 18

        # Fila de TOTALES
        n = len(self.filas)
        if n > 0:
            ft = n + 3          # fila total (1 título + 1 cabecera + n datos + 1)
            tf = PatternFill("solid", start_color=COLOR_TOTAL_BG)
            for c in range(1, n_cols + 1):
                cell        = ws.cell(row=ft, column=c)
                cell.fill   = tf
                cell.border = border
                cell.font   = bold

            ws.cell(row=ft, column=1).value     = "TOTAL"
            ws.cell(row=ft, column=1).alignment = Alignment(horizontal="left", vertical="center")

            for col_sum in (col_base, col_cuota, col_total):
                letra = get_column_letter(col_sum)
                cell  = ws.cell(row=ft, column=col_sum)
                cell.value         = f"=SUM({letra}3:{letra}{ft-1})"
                cell.number_format = '#,##0.00 €'
                cell.alignment     = Alignment(horizontal="right", vertical="center")

            cell = ws.cell(row=ft, column=col_n_fac)
            cell.value         = f"=COUNTA(A3:A{ft-1})"
            cell.number_format = '0 "facturas"'
            cell.alignment     = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[ft].height = 22

    def _hoja_resumen(self, ws):
        border = self._border()
        font   = Font(name="Arial", size=10)
        bold   = Font(name="Arial", size=10, bold=True)
        n_cols = len(COLUMNAS_RESUMEN)

        # Fila 1: Título
        ws.merge_cells(f"A1:{get_column_letter(n_cols)}1")
        t = ws["A1"]
        t.value     = "Resumen por Proveedor"
        t.font      = Font(name="Arial", size=12, bold=True, color="FFFFFF")
        t.fill      = PatternFill("solid", start_color="0D3B66")
        t.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        # Fila 2: Cabecera
        hdr_fill = PatternFill("solid", start_color=COLOR_HEADER_BG)
        hdr_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        for c, (campo, etiqueta, ancho) in enumerate(COLUMNAS_RESUMEN, 1):
            cell = ws.cell(row=2, column=c, value=etiqueta)
            cell.font      = hdr_font
            cell.fill      = hdr_fill
            cell.border    = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(c)].width = ancho
        ws.row_dimensions[2].height = 28

        # Agrupar por NIF
        proveedores: dict[str, dict] = {}
        for fila in self.filas:
            if fila.estado == "error":
                continue
            nif    = fila.datos.get("nif_emisor") or "SIN NIF"
            nombre = fila.datos.get("nombre_emisor") or "Desconocido"
            base   = float(fila.datos.get("base_imponible") or 0)
            cuota  = float(fila.datos.get("cuota_iva")     or 0)
            total  = float(fila.datos.get("total_factura") or 0)
            if nif not in proveedores:
                proveedores[nif] = {"nombre": nombre, "n": 0, "base": 0.0, "cuota": 0.0, "total": 0.0}
            p = proveedores[nif]
            p["n"] += 1; p["base"] += base; p["cuota"] += cuota; p["total"] += total

        provs = sorted(proveedores.items(), key=lambda x: x[1]["total"], reverse=True)

        fill_a = PatternFill("solid", start_color=COLOR_RESUMEN_BG)
        fill_b = PatternFill("solid", start_color="FFFFFF")

        for r, (nif, d) in enumerate(provs, 3):
            fill   = fill_a if r % 2 == 0 else fill_b
            valores = [nif, d["nombre"], d["n"], d["base"], d["cuota"], d["total"]]
            for c, (valor, (campo, _, _)) in enumerate(zip(valores, COLUMNAS_RESUMEN), 1):
                cell        = ws.cell(row=r, column=c, value=valor)
                cell.font   = font
                cell.fill   = fill
                cell.border = border
                if campo in ("base_imponible", "cuota_iva", "total_factura"):
                    cell.number_format = '#,##0.00 €'
                    cell.alignment     = Alignment(horizontal="right", vertical="center")
                elif campo == "_n_facturas":
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[r].height = 18

        # Fila de TOTALES
        if provs:
            ft = len(provs) + 3
            tf = PatternFill("solid", start_color=COLOR_TOTAL_BG)
            for c in range(1, n_cols + 1):
                cell        = ws.cell(row=ft, column=c)
                cell.fill   = tf
                cell.border = border
                cell.font   = bold

            ws.cell(row=ft, column=1).value     = "TOTAL"
            ws.cell(row=ft, column=1).alignment = Alignment(horizontal="left", vertical="center")

            for c, (campo, _, _) in enumerate(COLUMNAS_RESUMEN, 1):
                letra = get_column_letter(c)
                cell  = ws.cell(row=ft, column=c)
                if campo == "_n_facturas":
                    cell.value         = f"=SUM({letra}3:{letra}{ft-1})"
                    cell.number_format = '0'
                    cell.alignment     = Alignment(horizontal="center", vertical="center")
                elif campo in ("base_imponible", "cuota_iva", "total_factura"):
                    cell.value         = f"=SUM({letra}3:{letra}{ft-1})"
                    cell.number_format = '#,##0.00 €'
                    cell.alignment     = Alignment(horizontal="right", vertical="center")
            ws.row_dimensions[ft].height = 22

    # ── CSV ───────────────────────────────────

    def _construir_csv(self, ruta: Path):
        cabecera = [etiqueta for _, etiqueta, _ in COLUMNAS_EXCEL]
        with open(ruta, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(cabecera)
            for fila in self.filas:
                writer.writerow([fila.valor(campo) for campo, _, _ in COLUMNAS_EXCEL])

    # ── Consola ───────────────────────────────

    def resumen_consola(self) -> str:
        lineas = [
            "═" * 60,
            "  INFORME DE PROCESADO POR LOTES",
            "═" * 60,
            f"  Total archivos procesados : {self.total}",
            f"  ✅ Facturas válidas        : {self.validas}",
            f"  ⚠️  Con advertencias        : {self.con_advertencias}",
            f"  ❌ Errores                 : {self.con_error}",
            "─" * 60,
        ]
        if self.con_error:
            lineas.append("  Archivos con error:")
            for f in self.filas:
                if f.estado == "error":
                    lineas.append(f"    ✘ {f.archivo}: {f.error_log}")
        lineas.append("═" * 60)
        return "\n".join(lineas)


# ─────────────────────────────────────────────
# Procesador por lotes
# ─────────────────────────────────────────────

class ProcesadorLote:
    """
    Escanea carpetas o listas de archivos y genera un InformeLote.

    Uso:
        procesador = ProcesadorLote()
        informe = procesador.procesar_carpeta("/ruta/a/facturas/")
        informe.guardar_ambos("/home/juang/factura-ia/data/output/")
    """

    EXTENSIONES = EXTENSIONES_PDF | EXTENSIONES_IMAGEN

    def __init__(self, extractor: FacturaExtractor | None = None):
        self.extractor = extractor or FacturaExtractor()

    def procesar_carpeta(self, carpeta: str | Path, recursivo: bool = False) -> InformeLote:
        carpeta = Path(carpeta)
        if not carpeta.is_dir():
            raise NotADirectoryError(f"La carpeta no existe: {carpeta}")
        patron   = "**/*" if recursivo else "*"
        archivos = sorted(p for p in carpeta.glob(patron)
                          if p.is_file() and p.suffix.lower() in self.EXTENSIONES)
        if not archivos:
            logger.warning(f"No se encontraron facturas en: {carpeta}")
            return InformeLote()
        logger.info(f"📂 {carpeta}  |  {len(archivos)} archivos encontrados.")
        return self._procesar_lista(archivos)

    def procesar_archivos(self, rutas: list[str | Path]) -> InformeLote:
        return self._procesar_lista([Path(r) for r in rutas])

    def _procesar_lista(self, archivos: list[Path]) -> InformeLote:
        informe = InformeLote()
        total   = len(archivos)
        for idx, ruta in enumerate(archivos, 1):
            logger.info(f"[{idx}/{total}] {ruta.name}")
            fila = self._procesar_uno(ruta)
            informe.filas.append(fila)
            icono = {"válida": "✅", "incompleta": "⚠️ ", "error": "❌"}.get(fila.estado, "?")
            logger.info(f"  {icono} {fila.estado}  (confianza: {fila.confianza:.0%})")

        # ── Detección de duplicados ───────────────────────────────────────────
        vistos: dict[tuple, str] = {}
        for fila in informe.filas:
            if fila.estado == "error":
                continue
            nif    = fila.datos.get("nif_emisor")
            numero = fila.datos.get("numero_factura")
            if not nif or not numero:
                continue
            clave = (str(nif).upper().strip(), str(numero).upper().strip())
            if clave in vistos:
                aviso = f"POSIBLE FACTURA DUPLICADA (mismo NIF+numero que {vistos[clave]})"
                if aviso not in fila.advertencias:
                    fila.advertencias.append(aviso)
                fila.error_log = (fila.error_log + "; " + aviso).lstrip("; ")
            else:
                vistos[clave] = fila.archivo
        return informe

    def _procesar_uno(self, ruta: Path) -> FilaFactura:
        try:
            resultado: ResultadoExtraccion = self.extractor.extraer_archivo(ruta)
            estado = "válida" if resultado.es_valida() else "incompleta"
            error_log = "; ".join(resultado.advertencias) if resultado.advertencias else ""
            
            # Si es incompleta, añadir qué campos faltan al log
            if estado == "incompleta":
                from core.extractor import CAMPOS_REQUERIDOS
                vacios = [c for c in CAMPOS_REQUERIDOS if not resultado.datos.get(c)]
                if vacios:
                    nota_vacios = "Campos requeridos sin valor: " + ", ".join(vacios)
                    error_log = (error_log + "; " + nota_vacios).lstrip("; ")
            return FilaFactura(
                archivo=ruta.name, datos=resultado.datos,
                estado=estado, confianza=resultado.confianza,
                advertencias=resultado.advertencias, error_log=error_log,
            )
        except Exception as exc:
            logger.error(f"  ✘ {ruta.name}: {exc}")
            return FilaFactura(
                archivo=ruta.name, datos={},
                estado="error", confianza=0.0,
                advertencias=[], error_log=str(exc),
            )


# ─────────────────────────────────────────────
# Función de conveniencia
# ─────────────────────────────────────────────

def procesar_carpeta(
    carpeta: str | Path,
    output: str | Path = OUTPUT_DIR_DEFAULT,
    solo_excel: bool = False,
    solo_csv: bool = False,
    recursivo: bool = False,
) -> InformeLote:
    """Atajo rápido: escanea carpeta y guarda los resultados."""
    informe = ProcesadorLote().procesar_carpeta(carpeta, recursivo=recursivo)
    if solo_csv:       informe.guardar_csv(output)
    elif solo_excel:   informe.guardar_excel(output)
    else:              informe.guardar_ambos(output)
    return informe


# ─────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────

if __name__ == "__main__":
    import argparse

    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s", datefmt="%H:%M:%S")

    parser = argparse.ArgumentParser(description="FacturaIA — Procesador por lotes")
    parser.add_argument("entradas", nargs="+", help="Carpeta con facturas o lista de archivos PDF/JPG/PNG")
    parser.add_argument("--output", "-o", default=str(OUTPUT_DIR_DEFAULT), help="Directorio de salida")
    parser.add_argument("--solo-excel", action="store_true")
    parser.add_argument("--solo-csv",   action="store_true")
    parser.add_argument("--recursivo",  "-r", action="store_true", help="Buscar en subcarpetas")
    args = parser.parse_args()

    output_dir = Path(args.output)
    procesador = ProcesadorLote()
    entradas   = [Path(e) for e in args.entradas]

    if len(entradas) == 1 and entradas[0].is_dir():
        informe = procesador.procesar_carpeta(entradas[0], recursivo=args.recursivo)
    else:
        informe = procesador.procesar_archivos(entradas)

    print(informe.resumen_consola())
    if not informe.filas:
        sys.exit(0)

    output_dir.mkdir(parents=True, exist_ok=True)

    if args.solo_csv:
        print(f"\n📄 CSV   → {informe.guardar_csv(output_dir)}")
    elif args.solo_excel:
        print(f"\n📊 Excel → {informe.guardar_excel(output_dir)}")
    else:
        xlsx, csv_ = informe.guardar_ambos(output_dir)
        print(f"\n📊 Excel → {xlsx}")
        print(f"📄 CSV   → {csv_}")
